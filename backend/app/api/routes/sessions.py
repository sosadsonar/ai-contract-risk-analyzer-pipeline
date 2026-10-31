"""Session lifecycle HTTP contract; results are explicitly dev samples."""
import io
from uuid import UUID

from docx import Document
from fastapi import APIRouter, Request
from fastapi.responses import JSONResponse, Response

from app.api.errors import api_error
from app.api.upload import parse_upload
from app.core.enums import Role, SessionStatus
from app.session.store import Session, utcnow

router = APIRouter(prefix="/api/v1/sessions", tags=["sessions"])


def format_timestamp(value):
    return value.isoformat().replace("+00:00", "Z")


def status_payload(session: Session) -> dict:
    completed = {SessionStatus.UPLOADED: 0, SessionStatus.VALIDATING: 1,
                 SessionStatus.PROCESSING: 2, SessionStatus.COMPLETED: 3}.get(session.status, 0)
    return {"session_id": str(session.session_id), "status": session.status.value,
            "stage": session.stage.value, "manual_retry_count": session.manual_retry_count,
            "progress": {"completed_units": completed, "total_units": 3},
            "error_code": session.error_code, "updated_at": format_timestamp(session.updated_at)}


def lookup(request: Request, raw_id: str) -> Session | JSONResponse:
    try:
        sid = UUID(raw_id)
        if sid.version != 4 or str(sid) != raw_id.lower():
            raise ValueError("Invalid UUIDv4")
    except ValueError:
        return api_error(404, "SESSION_NOT_FOUND")
    request.app.state.cleanup.run_once()
    session = request.app.state.store.get(sid)
    if session is not None:
        return session
    return api_error(410 if request.app.state.db.tombstone(sid) else 404,
                     "SESSION_EXPIRED" if request.app.state.db.tombstone(sid) else "SESSION_NOT_FOUND")


def ready(session: Session) -> JSONResponse | None:
    if session.status == SessionStatus.COMPLETED:
        return None
    if session.status == SessionStatus.REJECTED:
        return api_error(409, "SESSION_REJECTED", details={"error_code": session.error_code})
    return api_error(409, "SESSION_NOT_READY")


@router.post("", status_code=201)
async def create(request: Request):
    state = request.app.state
    session = None
    rejection = None

    def on_role(value):
        nonlocal session, rejection
        if session is not None or rejection is not None:
            return
        try:
            role = Role(value)
        except ValueError:
            rejection = api_error(400, "ROLE_REQUIRED")
            return
        if not state.rate.check_and_record(request.client.host if request.client else "", utcnow()):
            rejection = api_error(429, "RATE_LIMITED")
        elif not state.quota.admit_new_session():
            rejection = api_error(429, "QUOTA_EXCEEDED")
        else:
            token = state.lease.try_acquire()
            if token is None:
                rejection = api_error(409, "SYSTEM_BUSY")
            else:
                try:
                    session = state.store.create(role)
                    state.lease.assign_owner(session.session_id)
                except Exception:
                    state.lease.release(token)
                    raise

    upload = await parse_upload(request, state.settings.MAX_UPLOAD_MB * 1024 * 1024, on_role)
    if rejection is not None:
        return rejection
    if session is None:
        return api_error(400, "ROLE_REQUIRED")
    # Dev sample selection stores only a fixed variant key, never the raw filename.
    filename = (upload.filename or "").lower()
    session.mock_variant = (
        "norisk" if "norisk" in filename else
        "insufficient" if "insufficient" in filename else None
    )
    state.orchestrator.start(session, upload.error)
    return JSONResponse(status_code=201, content={
        "session_id": str(session.session_id), "status": "UPLOADED",
        "created_at": format_timestamp(session.created_at)})


@router.get("/{session_id}/status")
def status(request: Request, session_id: str):
    session = lookup(request, session_id)
    if isinstance(session, JSONResponse):
        if session.status_code == 410:
            when = request.app.state.db.tombstone(UUID(session_id))
            return {"session_id": session_id, "status": "EXPIRED", "stage": None,
                    "manual_retry_count": 0, "progress": {"completed_units": 0, "total_units": 0},
                    "error_code": None, "updated_at": format_timestamp(when)}
        return session
    return status_payload(session)


@router.get("/{session_id}/result")
def result(request: Request, session_id: str):
    session = lookup(request, session_id)
    if isinstance(session, JSONResponse):
        return session
    request.app.state.store.touch(session.session_id)
    if error := ready(session):
        return error
    return session.result


@router.post("/{session_id}/retry")
async def retry(request: Request, session_id: str):
    session = lookup(request, session_id)
    if isinstance(session, JSONResponse):
        return session
    request.app.state.store.touch(session.session_id)
    if session.status != SessionStatus.FAILED:
        return api_error(409, "INVALID_SESSION_STATE", details={"status": session.status.value})
    if session.manual_retry_count >= 2:
        return api_error(409, "RETRY_LIMIT_EXCEEDED")
    if not request.app.state.orchestrator.retry(session):
        return api_error(409, "SYSTEM_BUSY")
    return status_payload(session)


@router.get("/{session_id}/export")
def export(request: Request, session_id: str, format: str = "docx"):
    session = lookup(request, session_id)
    if isinstance(session, JSONResponse):
        return session
    request.app.state.store.touch(session.session_id)
    if error := ready(session):
        return error
    if format != "docx":
        return api_error(501, "FEATURE_NOT_ENABLED")
    result = session.result
    report = Document()
    report.add_heading("Phân tích rủi ro hợp đồng lao động", 0)
    for clause in result["clauses"]:
        report.add_heading(clause["clause_id"], level=1)
        report.add_paragraph(clause["original_text"])
        for risk in result["risks"]:
            if risk["clause_id"] == clause["clause_id"]:
                report.add_paragraph(f"Rủi ro: {', '.join(risk['risk_labels'])} ({risk['severity']})")
                report.add_paragraph(risk["explanation"])
                report.add_paragraph(f"Căn cứ: {', '.join(risk['citation_ids'])}")
    for citation in result["citations"]:
        report.add_paragraph(f"{citation['citation_id']}: {citation['content']}")
    report.add_paragraph(result["disclaimer"])
    report.add_paragraph(result["data_notice"])
    buffer = io.BytesIO()
    report.save(buffer)
    return Response(buffer.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={"Content-Disposition": 'attachment; filename="report.docx"'})


@router.delete("/{session_id}")
def end(request: Request, session_id: str):
    session = lookup(request, session_id)
    if isinstance(session, JSONResponse) and session.status_code != 410:
        return session
    if not isinstance(session, JSONResponse):
        request.app.state.store.touch(session.session_id)
        request.app.state.orchestrator.end(session.session_id)
    return {"session_id": session_id, "status": "EXPIRED",
            "message": "Phiên đã kết thúc và dữ liệu phiên đã được xóa."}
